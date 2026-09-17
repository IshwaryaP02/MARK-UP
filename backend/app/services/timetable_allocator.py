"""
Timetable Allocator Service
===========================
Reads teacher–subject–hours requirements from an Excel file and generates
a weekly timetable that satisfies all hard constraints via backtracking.

Hard Constraints
----------------
1. Shifts: 2 shifts/day, each with 5 periods of 50 minutes. Students in
   Shift 1 are different from Shift 2 — completely isolated.
2. Shift Isolation: A teacher works EXCLUSIVELY in one shift.
3. No Teacher Collision: A teacher can teach at most one class at a time.
4. No Class Collision: A class has at most one teacher per period.
5. Daily Limit: A subject cannot repeat on the same day for the same class.
6. 1-to-1 Subject Mapping: One subject → one teacher per class.
7. Hour Fulfillment: Total allocated hours must match Required_Hours exactly.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

DAYS = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday")
PERIODS = tuple(range(1, 6))          # 5 periods per shift
SHIFTS = ("First Shift", "Second Shift")
MAX_SEARCH_NODES = 500_000             # Bail-out to prevent infinite backtracking


@dataclass(frozen=True)
class Requirement:
    teacher: str          # display name (from Excel)
    employee_id: str      # Faculty_ID column → maps to users.employee_id
    subject: str          # subject code / short name
    subject_name: str     # full subject name (if provided)
    hours: int            # weekly hours required
    shift: str            # "First Shift" | "Second Shift"


def _key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _find_column(headers: dict[str, str], names: tuple[str, ...]) -> str | None:
    for name in names:
        wanted = _key(name)
        for actual, normalized in headers.items():
            if normalized == wanted or wanted in normalized or normalized in wanted:
                return actual
    return None


def _normalize_shift(value: Any) -> str:
    v = _key(value)
    if v in {"1", "first", "firstshift", "morning", "shift1", "s1", "i"}:
        return SHIFTS[0]
    if v in {"2", "second", "secondshift", "evening", "afternoon", "shift2", "s2", "ii"}:
        return SHIFTS[1]
    raise ValueError(
        f"Cannot interpret shift '{value}'. Use 'First Shift', 'Second Shift', 1, or 2."
    )


def read_requirements(content: bytes) -> list[Requirement]:
    """Parse the Excel file and return a list of Requirement objects."""
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    try:
        raw_headers = next(rows)
    except StopIteration as exc:
        raise ValueError("The Excel file is empty.") from exc

    headers = {str(v).strip(): _key(v) for v in raw_headers if v is not None}

    teacher_col  = _find_column(headers, ("Teacher_Name", "Teacher", "Faculty", "Faculty_Name", "Instructor", "Name"))
    emp_id_col   = _find_column(headers, ("Faculty_ID", "Employee_ID", "Emp_ID", "FacultyID", "EmployeeID", "EmpID", "ID"))
    subject_col  = _find_column(headers, ("Subject", "Subject_Code", "Course", "Subject_Name", "Code"))
    sub_name_col = _find_column(headers, ("Subject_Full_Name", "Full_Name", "SubjectName", "Course_Name"))
    hours_col    = _find_column(headers, ("Required_Hours", "Hours", "Weekly_Hours", "Periods", "Required", "Hrs"))
    shift_col    = _find_column(headers, ("Shift_Assigned", "Shift", "Session", "Shift_No"))

    missing = [
        name for name, col in [
            ("Teacher_Name", teacher_col),
            ("Subject", subject_col),
            ("Required_Hours", hours_col),
            ("Shift_Assigned", shift_col),
        ] if not col
    ]
    if missing:
        raise ValueError(f"Missing required Excel columns: {', '.join(missing)}")

    col_index = {name: list(headers).index(name) for name in headers}

    requirements: list[Requirement] = []
    for row_num, row in enumerate(rows, start=2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        teacher   = str(row[col_index[teacher_col]] or "").strip()
        subject   = str(row[col_index[subject_col]] or "").strip()
        sub_name  = str(row[col_index[sub_name_col]] or subject).strip() if sub_name_col else subject
        emp_id    = str(row[col_index[emp_id_col]] or "").strip() if emp_id_col else ""
        shift_raw = row[col_index[shift_col]]
        try:
            hours = int(float(row[col_index[hours_col]]))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Row {row_num}: Required_Hours must be a number.") from exc

        if not teacher or not subject or hours < 1:
            raise ValueError(f"Row {row_num}: teacher, subject, and hours >= 1 are required.")

        requirements.append(
            Requirement(
                teacher=teacher,
                employee_id=emp_id,
                subject=subject,
                subject_name=sub_name,
                hours=hours,
                shift=_normalize_shift(shift_raw),
            )
        )

    if not requirements:
        raise ValueError("No valid rows found in the Excel file.")

    # Validate: each teacher must have a consistent shift
    teacher_shifts: dict[str, str] = {}
    for req in requirements:
        prev = teacher_shifts.setdefault(req.teacher, req.shift)
        if prev != req.shift:
            raise ValueError(
                f"Teacher '{req.teacher}' is assigned to both shifts — each teacher must be in exactly one shift."
            )

    return requirements


def generate_timetable(
    requirements: list[Requirement],
    sections: list[str],
) -> list[dict[str, Any]]:
    """
    Generate a collision-free weekly timetable via constraint-satisfaction
    backtracking.

    Returns a list of slot dicts:
      {day, period, section, shift, teacher, employee_id, subject, subject_name}
    """
    if not sections:
        raise ValueError("At least one section is required.")

    # ── Pre-process ──────────────────────────────────────────────────────────
    # Sort: most constrained (most hours) first for better pruning
    work = sorted(requirements, key=lambda r: r.hours, reverse=True)

    # Map sections to their shift (caller should split sections by shift)
    # We treat all sections as equal-opportunity and rely on shift isolation.

    # ── Constraint state ─────────────────────────────────────────────────────
    # teacher_busy[(teacher, day, period)] → True when occupied
    teacher_busy: set[tuple[str, str, int]] = set()
    # class_busy[(section, day, period, shift)] → True when occupied
    class_busy: set[tuple[str, str, int, str]] = set()
    # subject_days[(section, day, subject)] → prevents repeat on same day
    subject_days: set[tuple[str, str, str]] = set()
    # subject_teacher[(section, subject)] → enforces 1-to-1 mapping
    subject_teacher: dict[tuple[str, str], str] = {}

    allocation: list[dict[str, Any]] = []

    # All possible (day, period, section, shift) combos — ordered for determinism
    all_slots = [
        (day, period, section, shift)
        for day in DAYS
        for shift in SHIFTS
        for period in PERIODS
        for section in sections
    ]

    nodes_visited = [0]

    def place(req_idx: int, remaining: int) -> bool:
        nodes_visited[0] += 1
        if nodes_visited[0] > MAX_SEARCH_NODES:
            raise ValueError(
                "Timetable solver exceeded search limit. "
                "Reduce required hours or add more sections/days."
            )

        if req_idx == len(work):
            return remaining == 0

        req = work[req_idx]

        if remaining == 0:
            # Move to next requirement
            next_idx = req_idx + 1
            if next_idx >= len(work):
                return True
            return place(next_idx, work[next_idx].hours)

        for day, period, section, shift in all_slots:
            # Shift isolation
            if shift != req.shift:
                continue

            teacher_key    = (req.teacher, day, period)
            class_key      = (section, day, period, shift)
            day_subj_key   = (section, day, req.subject)
            mapping_key    = (section, req.subject)

            if teacher_key  in teacher_busy:  continue
            if class_key    in class_busy:    continue
            if day_subj_key in subject_days:  continue

            # 1-to-1 subject-teacher per class
            mapped = subject_teacher.get(mapping_key)
            if mapped and mapped != req.teacher:
                continue

            # Place the slot
            teacher_busy.add(teacher_key)
            class_busy.add(class_key)
            subject_days.add(day_subj_key)
            subject_teacher[mapping_key] = req.teacher
            allocation.append({
                "day": day,
                "period": period,
                "section": section,
                "shift": shift,
                "teacher": req.teacher,
                "employee_id": req.employee_id,
                "subject": req.subject,
                "subject_name": req.subject_name,
            })

            if place(req_idx, remaining - 1):
                return True

            # Backtrack
            allocation.pop()
            teacher_busy.discard(teacher_key)
            class_busy.discard(class_key)
            subject_days.discard(day_subj_key)
            # Only remove subject_teacher mapping if no other slot still uses it
            if not any(
                s["section"] == section and s["subject"] == req.subject
                for s in allocation
            ):
                subject_teacher.pop(mapping_key, None)

        return False

    if not place(0, work[0].hours):
        raise ValueError(
            "No feasible timetable found. "
            "Check that required hours fit within the 6-day × 5-period grid "
            "without violating shift, collision, or daily-repeat constraints."
        )

    return allocation
